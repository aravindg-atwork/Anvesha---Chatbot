"""Bounded Office extraction with exact sheet/slide locators; never executes macros."""
import io
from pathlib import Path
from zipfile import ZipFile, BadZipFile


def check_package(path: Path):
    try:
        with ZipFile(path) as package:
            if len(package.infolist()) > 10000 or sum(item.file_size for item in package.infolist()) > 100 * 1024 * 1024:
                raise ValueError('This Office document expands beyond the 100 MB processing limit.')
    except BadZipFile:
        raise ValueError('This Office document is damaged or encrypted. Save an unencrypted copy and retry.') from None


def spreadsheet(path: Path, report):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    check_package(path)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    pages, chars, cells = [], 0, 0
    try:
        if len(workbook.worksheets) > 100:
            raise ValueError('Workbooks are limited to 100 sheets.')
        for sheet in workbook.worksheets:
            # Honor actual content even if a third-party writer stored incorrect dimensions.
            sheet.reset_dimensions()
            headers = None
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if number > 50000 or len(row) > 200:
                    raise ValueError('Sheets are limited to 50,000 rows and 200 columns.')
                cells += len(row)
                if cells > 500000:
                    raise ValueError('Workbooks are limited to 500,000 cells.')
                values = [(i, str(value).strip()) for i, value in enumerate(row) if value is not None and str(value).strip()]
                if not values:
                    continue
                if headers is None:
                    headers = {i: value for i, value in values}
                    text = f'Sheet: {sheet.title}. Header row {number}: ' + ' | '.join(f'{get_column_letter(i+1)}: {value}' for i, value in values)
                else:
                    text = f'Sheet: {sheet.title}. Row {number}: ' + ' | '.join(f'{headers.get(i, get_column_letter(i+1))} ({get_column_letter(i+1)}{number}): {value}' for i, value in values)
                chars += len(text)
                if chars > 3_000_000:
                    raise ValueError('Extracted text exceeds the 3 MB limit. Split this workbook.')
                pages.append((text, {'kind': 'sheet', 'sheet': sheet.title, 'range': f'A{number}:{get_column_letter(max(i for i, _ in values)+1)}{number}'}))
            report('Workbook reader', 40, f'Read sheet {sheet.title}.')
    finally:
        workbook.close()
    return pages


def presentation(path: Path, report, ocr):
    from pptx import Presentation
    from PIL import Image
    check_package(path)
    deck = Presentation(path)
    if len(deck.slides) > 250:
        raise ValueError('Presentations are limited to 250 slides.')
    pages, chars, images = [], 0, 0

    def read_shapes(shapes):
        nonlocal images
        parts = []
        for shape in shapes:
            if shape.shape_type == 6:  # Group: recurse in document order.
                parts.extend(read_shapes(shape.shapes))
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                parts.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
            if shape.shape_type == 13:  # Picture: OCR embedded image only.
                images += 1
                if images > 100:
                    raise ValueError('Presentations are limited to 100 embedded images for OCR.')
                with Image.open(io.BytesIO(shape.image.blob)) as image:
                    parts.append(ocr(image))
        return parts

    for i, slide in enumerate(deck.slides, 1):
        lines = read_shapes(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes and notes.text.strip():
                lines.append('Speaker notes: ' + notes.text)
        text = '\n'.join(part for part in lines if part.strip())
        chars += len(text)
        if chars > 3_000_000:
            raise ValueError('Extracted text exceeds the 3 MB limit. Split this presentation.')
        pages.append((text, {'kind': 'slide', 'slide': i}))
        report('Slide reader', 10 + int(45 * i / max(1, len(deck.slides))), f'Read slide {i} of {len(deck.slides)}.')
    return pages, images
