"""Real Office extraction plus mocked Google transport, account and sync boundaries."""
import io
import json
import os
import time
import unittest
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlparse

import httpx
from fastapi import HTTPException
from backend import test_product as product
from backend.db import connect, project_data
from backend.drive import encrypt, run_sync, SCOPE, FOLDER, validate_passage_access
from backend.extraction import extract


def workbook_bytes():
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Products'
    sheet.append(['Product', 'Warranty months'])
    sheet.append(['Atlas', 24])
    other = workbook.create_sheet('Support')
    other.append(['Contact', 'Email'])
    other.append(['Help desk', 'help@example.com'])
    result = io.BytesIO(); workbook.save(result)
    return result.getvalue()


def slides_bytes():
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Atlas setup'
    slide.placeholders[1].text = 'Connect the supplied power cable.'
    slide.notes_slide.notes_text_frame.text = 'Use the adapter supplied in the box.'
    result = io.BytesIO(); deck.save(result)
    return result.getvalue()


class ImportTests(unittest.TestCase):
    save = product.ProductTests.save
    upload = product.ProductTests.upload
    run_job_id = product.ProductTests.run_job_id

    def setUp(self):
        product.ProductTests.setUp(self)

    def connect_drive(self):
        self.settings = patch.dict(os.environ, {'GOOGLE_OAUTH_CLIENT_ID': 'test-client', 'GOOGLE_OAUTH_CLIENT_SECRET': 'test-secret',
            'GOOGLE_OAUTH_REDIRECT_URI': 'http://testserver/api/drive/callback'})
        self.settings.start(); self.addCleanup(self.settings.stop)
        with connect() as db:
            user_id = db.execute('SELECT id FROM users').fetchone()['id']
            db.execute('INSERT INTO drive_connections VALUES (?,?,?,1)', (user_id, encrypt({'access_token': 'access-token-secret', 'refresh_token': 'refresh-token-secret', 'expires': time.time()+3600}), 'alice@google.test'))
        return user_id

    def test_xlsx_and_pptx_uploads_prepare_exact_locators(self):
        for name, content in [('catalog.xlsx', workbook_bytes()), ('setup.pptx', slides_bytes())]:
            response = self.client.post('/api/projects/project-one/sources', files={'file': (name, content)})
            self.assertEqual(response.status_code, 200, response.text)
            self.run_job_id(response.json()['jobId'])
        sources = self.client.get('/api/state').json()['state']['assistants'][0]['draft']['sources']
        sheet, deck = sorted(sources, key=lambda source: source['fileType'], reverse=True)
        self.assertEqual(sheet['fileType'], 'XLSX')
        self.assertEqual(sheet['status'], 'ready')
        self.assertEqual({c['locator']['sheet'] for c in sheet['chunks']}, {'Products', 'Support'})
        row = next(c for c in sheet['chunks'] if 'Atlas' in c['text'])
        self.assertEqual(row['locator']['range'], 'A2:B2')
        self.assertIn('Warranty months (B2): 24', row['text'])
        self.assertEqual(deck['fileType'], 'PPTX')
        self.assertEqual(deck['chunks'][0]['locator'], {'kind': 'slide', 'slide': 1})
        self.assertIn('supplied power cable', deck['chunks'][0]['text'])
        self.assertIn('adapter supplied', deck['chunks'][0]['text'])

    def test_media_requires_consent_and_uses_timestamped_transcript(self):
        endpoint = '/api/projects/project-one/sources'
        self.assertEqual(self.client.post(endpoint, files={'file': ('meeting.mp3', b'audio')}).status_code, 400)
        self.assertEqual(self.client.post(endpoint, files={'file': ('meeting.mp3', b'audio')}, data={'mediaConsent': 'true'}).status_code, 503)
        with patch.dict(os.environ, {'GOOGLE_API_KEY': 'fake-test-key'}):
            response = self.client.post(endpoint, files={'file': ('meeting.mp3', b'audio')}, data={'mediaConsent': 'true'})
            self.assertEqual(response.status_code, 200, response.text)
            with patch('backend.media.transcribe', return_value=[('The launch is on Friday.', {'kind': 'media', 'startSeconds': 12, 'endSeconds': 18})]):
                self.run_job_id(response.json()['jobId'])
        source = self.client.get('/api/state').json()['state']['assistants'][0]['draft']['sources'][0]
        self.assertEqual(source['chunks'][0]['locator']['startSeconds'], 12)
        self.assertIn('Gemini', source['extractionNote'])

    def test_media_sdk_response_validation_and_client_cleanup(self):
        from backend.media import transcribe
        from backend.db import data_dir
        from unittest.mock import MagicMock
        media = data_dir() / 'test.wav'; media.write_bytes(b'fake-audio-for-mocked-transport')
        client = MagicMock()
        client.models.generate_content.return_value = SimpleNamespace(text=json.dumps({'segments': [
            {'startSeconds': 4, 'endSeconds': 7, 'text': 'The service opens at nine.'}]}))
        with patch.dict(os.environ, {'GOOGLE_API_KEY': 'fake-key'}), patch('google.genai.Client', return_value=client):
            pages = transcribe(media, lambda *args: None)
            self.assertEqual(pages[0][1], {'kind': 'media', 'startSeconds': 4.0, 'endSeconds': 7.0})
            self.assertEqual(client.models.generate_content.call_args.kwargs['contents'][0].inline_data.mime_type, 'audio/wav')
            client.close.assert_called_once()
            client.models.generate_content.return_value = SimpleNamespace(text=json.dumps({'segments': [
                {'startSeconds': 9, 'endSeconds': 1, 'text': 'Invalid timing'}]}))
            with self.assertRaisesRegex(ValueError, 'invalid timestamps'):
                transcribe(media, lambda *args: None)

    def test_oauth_state_pkce_replay_and_encrypted_tokens(self):
        self.connect_drive()
        result = self.client.post('/api/drive/connect', json={'projectId': 'project-one'})
        self.assertEqual(result.status_code, 200, result.text)
        query = parse_qs(urlparse(result.json()['url']).query)
        self.assertEqual(query['scope'], [SCOPE])
        self.assertEqual(query['code_challenge_method'], ['S256'])
        state = query['state'][0]
        self.assertEqual(self.client.get('/api/drive/callback', params={'state': 'wrong', 'code': 'a'}).status_code, 400)
        def provider(method, url, **kwargs):
            if url.endswith('/token'):
                self.assertTrue(kwargs['data']['code_verifier'])
                return httpx.Response(200, json={'access_token': 'new-token', 'refresh_token': 'new-refresh', 'scope': SCOPE, 'expires_in': 3600})
            return httpx.Response(200, json={'user': {'emailAddress': 'alice@google.test'}})
        with patch('backend.drive.provider_request', side_effect=provider):
            response = self.client.get('/api/drive/callback', params={'state': state, 'code': 'sample-code'}, follow_redirects=False)
        self.assertEqual(response.status_code, 303, response.text)
        with connect() as db:
            stored = db.execute('SELECT credentials FROM drive_connections').fetchone()['credentials']
        self.assertNotIn('new-token', stored)
        self.assertNotIn('new-refresh', stored)
        self.assertEqual(self.client.get('/api/drive/callback', params={'state': state, 'code': 'sample-code'}).status_code, 400)

    def test_drive_preview_sync_dedup_updates_and_removal(self):
        user_id = self.connect_drive()
        root = {'id': 'folder-one', 'name': 'Product guides', 'mimeType': FOLDER, 'parents': ['root']}
        file = {'id': 'file-one', 'name': 'Products', 'mimeType': 'application/vnd.google-apps.spreadsheet', 'version': '1', 'parents': ['folder-one']}
        unsupported = {'id': 'file-two', 'name': 'archive.zip', 'mimeType': 'application/zip', 'version': '1', 'parents': ['folder-one']}
        config = {'projectId': 'project-one', 'folderIds': ['folder-one'], 'destination': 'Business', 'recursive': True}
        with patch('backend.drive.metadata', return_value=root), patch('backend.drive.children', return_value={'files': [file, unsupported]}), patch('backend.drive.download', return_value=workbook_bytes()):
            preview = self.client.post('/api/drive/preview', json=config)
            self.assertEqual(preview.status_code, 200, preview.text)
            self.assertEqual(preview.json()['supported'], 1)
            self.assertEqual(preview.json()['types'], {'XLSX': 1, 'Other': 1})
            response = self.client.post('/api/drive/import', json=config)
            self.assertEqual(response.status_code, 200, response.text)
            job_id = response.json()['jobId']
            self.run_job_id(job_id)
            with connect() as db:
                ingest = db.execute("SELECT id FROM jobs WHERE kind='ingest'").fetchone()['id']
            self.run_job_id(ingest)
            self.run_job_id(job_id)
        sources = self.client.get('/api/state').json()['state']['assistants'][0]['draft']['sources']
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]['folderPath'], 'Business/Product guides')
        self.assertEqual(sources[0]['driveFileId'], 'file-one')
        with connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) AS n FROM jobs WHERE kind='ingest'").fetchone()['n'], 1)
        # A changed Drive revision is blocked at answer time until synced.
        project = {'id': 'project-one', 'draft': {'sources': sources}}
        with patch('backend.drive.metadata', side_effect=lambda _, fid: {**file, 'version': '2'} if fid == 'file-one' else root):
            with self.assertRaises(HTTPException) as raised:
                validate_passage_access(project, user_id, [{'sourceId': sources[0]['id']}])
            self.assertEqual(raised.exception.status_code, 409)
        file['version'] = '2'
        with patch('backend.drive.metadata', return_value=root), patch('backend.drive.children', return_value={'files': [file, unsupported]}), patch('backend.drive.download', return_value=workbook_bytes()):
            self.run_job_id(job_id)
            with connect() as db:
                latest = db.execute("SELECT id FROM jobs WHERE kind='ingest' ORDER BY rowid DESC LIMIT 1").fetchone()['id']
            self.run_job_id(latest)
        current = self.client.get('/api/state').json()['state']['assistants'][0]['draft']['sources']
        self.assertEqual(len(current), 1)
        self.assertEqual(current[0]['id'], sources[0]['id'])
        self.assertEqual(current[0]['driveRevision'], '2')
        with patch('backend.drive.metadata', return_value=root), patch('backend.drive.children', return_value={'files': []}):
            self.run_job_id(job_id)
        self.assertEqual(self.client.get('/api/state').json()['state']['assistants'][0]['draft']['sources'], [])

    def test_drive_revoked_folder_membership_and_disconnect_disable_sources(self):
        user_id = self.connect_drive()
        config = {'projectId': 'project-one', 'folderIds': ['selected-folder']}
        upload = self.upload()
        source = upload['source']; source.update(driveFileId='file-one', driveRevision='1')
        with connect() as db:
            db.execute('UPDATE sources SET data=? WHERE id=?', (json.dumps(source), source['id']))
            db.execute('INSERT INTO drive_items VALUES (?,?,?,?)', ('project-one', 'file-one', source['id'], 'fingerprint'))
            db.execute('INSERT INTO drive_links(project_id,user_id,config) VALUES (?,?,?)', ('project-one', user_id, json.dumps(config)))
        file = {'id': 'file-one', 'version': '1', 'parents': ['other-folder']}
        other = {'id': 'other-folder', 'mimeType': FOLDER, 'parents': []}
        with patch('backend.drive.metadata', side_effect=lambda _, fid: file if fid == 'file-one' else other):
            with self.assertRaises(HTTPException):
                validate_passage_access({'id': 'project-one', 'draft': {'sources': [source]}}, user_id, [{'sourceId': source['id']}])
        self.assertEqual(self.client.get('/api/state').json()['state']['assistants'][0]['draft']['sources'], [])
        with patch('backend.drive.provider_request', return_value=httpx.Response(200)):
            self.assertEqual(self.client.delete('/api/drive/connection').status_code, 200)
        with connect() as db:
            self.assertEqual(db.execute('SELECT credentials FROM drive_connections').fetchone()['credentials'], '')

    def test_drive_ownership_and_invalid_destination(self):
        self.connect_drive()
        self.assertEqual(self.client.post('/api/drive/preview', json={'projectId': 'project-one', 'folderIds': ['folder'], 'destination': '../escape'}).status_code, 400)
        self.client.post('/api/auth/signout')
        self.client.post('/api/auth/signup', json={'email': 'bob@example.com', 'name': 'Bob', 'password': 'correct-password'})
        self.assertEqual(self.client.get('/api/drive/status?projectId=project-one').status_code, 404)
        self.assertEqual(self.client.post('/api/drive/preview', json={'projectId': 'project-one', 'folderIds': ['folder']}).status_code, 404)

    def test_removing_a_drive_source_persists_its_sync_exclusion(self):
        user_id = self.connect_drive()
        source = self.upload()['source']; source['driveFileId'] = 'file-one'
        with connect() as db:
            db.execute('UPDATE sources SET data=? WHERE id=?', (json.dumps(source), source['id']))
            db.execute('INSERT INTO drive_items VALUES (?,?,?,?)', ('project-one', 'file-one', source['id'], 'f'))
            db.execute('INSERT INTO drive_links(project_id,user_id,config) VALUES (?,?,?)', ('project-one', user_id, json.dumps({'projectId': 'project-one', 'folderIds': ['folder-one']})))
        self.assertEqual(self.client.delete(f"/api/projects/project-one/sources/{source['id']}").status_code, 200)
        with connect() as db:
            config = json.loads(db.execute('SELECT config FROM drive_links').fetchone()['config'])
        self.assertEqual(config['excludedIds'], ['file-one'])

    def test_oauth_callback_rejects_a_revoked_anvesha_session(self):
        self.connect_drive()
        response = self.client.post('/api/drive/connect', json={'projectId': 'project-one'})
        state = parse_qs(urlparse(response.json()['url']).query)['state'][0]
        self.client.post('/api/auth/signout')
        self.assertEqual(self.client.get('/api/drive/callback', params={'state': state, 'code': 'not-used'}).status_code, 400)

    def test_registered_file_update_relabels_and_automatically_reindexes(self):
        from backend.source_watch import scan_sources
        from backend.vectors import signature
        created = self.upload('reference.txt', b'Product warranty covers the supplied adapter for 24 months.')
        self.run_job_id(created['jobId'])
        with connect() as db:
            row = db.execute('SELECT * FROM sources WHERE id=?', (created['source']['id'],)).fetchone()
            current = json.loads(row['data'])
            self.assertIn('Warranty', current['autoTags'])
            self.assertTrue(any('warranty' in question.lower() for question in current['suggestedQuestions']))
            current['tags'] = ['My reference']
            db.execute('UPDATE sources SET data=? WHERE id=?', (json.dumps(current), row['id']))
            project = project_data(db, db.execute('SELECT * FROM projects WHERE id=?', ('project-one',)).fetchone())
            db.execute('UPDATE projects SET index_revision=? WHERE id=?', (signature(project), 'project-one'))
            db.execute("UPDATE jobs SET status='complete' WHERE kind='index'")
            stored_path = Path(row['path'])
        scan_sources()  # Establish the current original's filesystem stamp.
        stored_path.write_text('Invoice payment is due within 30 days.', encoding='utf-8')
        scan_sources()
        with connect() as db:
            ingest = db.execute("SELECT * FROM jobs WHERE kind='ingest' ORDER BY rowid DESC LIMIT 1").fetchone()
            self.assertNotEqual(ingest['id'], created['jobId'])
            self.assertEqual(ingest['status'], 'queued')
        self.run_job_id(ingest['id'])
        with connect() as db:
            refreshed = json.loads(db.execute('SELECT data FROM sources WHERE id=?', (created['source']['id'],)).fetchone()['data'])
            self.assertIn('Payment', refreshed['autoTags'])
            self.assertNotIn('Warranty', refreshed['autoTags'])
            self.assertEqual(refreshed['tags'], ['My reference'])
            self.assertFalse(any('warranty' in question.lower() for question in refreshed['suggestedQuestions']))
            self.assertEqual(db.execute("SELECT COUNT(*) AS n FROM jobs WHERE kind='index' AND status='queued'").fetchone()['n'], 1)
        scan_sources()
        with connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) AS n FROM jobs WHERE kind='ingest'").fetchone()['n'], 2)
