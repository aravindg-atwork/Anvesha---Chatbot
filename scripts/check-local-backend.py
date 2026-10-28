"""Exercise real Docker OCR, durable jobs and Qdrant (no model/API mocks)."""
import io
import os
import time
import sys
import subprocess
from pathlib import Path
from uuid import uuid4

import pymupdf
import requests
from PIL import Image, ImageDraw, ImageFont
from docx import Document

BASE = os.getenv('ANVESHA_TEST_URL', 'http://127.0.0.1:8000')
client = requests.Session()


def call(method, path, **kwargs):
    response = client.request(method, BASE + path, timeout=90, **kwargs)
    assert response.ok, f'{method} {path}: {response.status_code} {response.text}'
    return response.json()


def wait_ready(project_id, kind):
    deadline = time.monotonic() + 360
    previous = None
    while time.monotonic() < deadline:
        result = call('GET', f'/api/projects/{project_id}/processing')
        jobs = [j for j in result['jobs'] if j['kind'] == kind]
        status = tuple((j['stage'], j['progress'], j['status']) for j in jobs)
        if status != previous:
            print(f'{kind}: {status}', flush=True)
            previous = status
        failures = [j for j in jobs if j['status'] == 'failed']
        assert not failures, str(failures)
        if jobs and all(j['status'] == 'complete' for j in jobs):
            return result
        time.sleep(1)
    raise AssertionError('Background jobs did not finish within six minutes.')


def main():
    health = call('GET', '/api/health')
    assert health['database'] and health['qdrant'], health
    identity = uuid4().hex[:10]
    email = f'smoke-{identity}@example.com'
    password = uuid4().hex
    call('POST', '/api/auth/signup', json={'email': email, 'password': password, 'name': 'Backend Smoke Test'})
    snapshot = call('GET', '/api/state')
    org_id, project_id = 'org-' + identity, 'project-' + identity
    snapshot['state']['orgs'] = [{'id': org_id, 'name': 'Smoke test workspace', 'isSample': False}]
    snapshot['state']['assistants'] = [{'id': project_id, 'orgId': org_id, 'draft': {
        'name': 'Smoke Test Assistant', 'purpose': 'Answer policy questions', 'sources': [], 'faqs': [], 'conflicts': []}}]
    call('PUT', '/api/state', json=snapshot)
    prefix = f'/api/projects/{project_id}'
    call('POST', prefix + '/sources', files={'file': ('leave.txt', b'Employees receive eighteen days of paid annual leave per year.')})

    image = Image.new('RGB', (1800, 400), 'white')
    font_path = next((p for p in [Path('C:/Windows/Fonts/arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')] if p.exists()), None)
    font = ImageFont.truetype(str(font_path), 52) if font_path else ImageFont.load_default(size=52)
    ImageDraw.Draw(image).multiline_text((50, 60), 'The office opens at 9 am.\nThe office closes at 6 pm.\nContact the People team for help.', font=font, fill='black', spacing=25)
    png = io.BytesIO()
    image.save(png, 'PNG')
    call('POST', prefix + '/sources', files={'file': ('office-notice.png', png.getvalue(), 'image/png')})
    pdf = pymupdf.open()
    page = pdf.new_page(width=900, height=200)
    page.insert_image(page.rect, stream=png.getvalue())
    call('POST', prefix + '/sources', files={'file': ('scanned-notice.pdf', pdf.tobytes(), 'application/pdf')})
    pdf.close()
    document = Document()
    document.add_heading('Travel policy', 1)
    document.add_paragraph('Employees may claim train tickets for approved business travel.')
    docx = io.BytesIO()
    document.save(docx)
    call('POST', prefix + '/sources', files={'file': ('travel.docx', docx.getvalue())})
    result = wait_ready(project_id, 'ingest')
    assert len(result['sources']) == 4
    for source in result['sources']:
        assert source['status'] == 'ready', source
        if source['fileType'] in ('Image', 'PDF'):
            assert '9 am' in ' '.join(c['text'] for c in source['chunks']), source
    print('Real PNG and scanned-PDF OCR, TXT and DOCX parsing passed.', flush=True)
    call('POST', prefix + '/train')
    result = wait_ready(project_id, 'index')
    assert result['indexed']
    response = call('POST', prefix + '/answer', json={'question': 'When does the office open?', 'engine': 'demo'})
    assert response['citations'] and '9 am' in response['text'], response
    if '--restart' in sys.argv:
        subprocess.run(['docker', 'compose', 'restart', 'app'], check=True, capture_output=True)
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            try:
                if client.get(BASE + '/api/health', timeout=5).ok:
                    break
            except requests.RequestException:
                pass
            time.sleep(1)
        response = call('POST', prefix + '/answer', json={'question': 'When does the office open?', 'engine': 'demo'})
        assert response['citations'] and '9 am' in response['text']
        print('Session, SQLite, uploads and Qdrant retrieval survived a Docker app restart.', flush=True)
    other = requests.Session()
    assert other.get(BASE + prefix + '/processing', timeout=30).status_code == 401
    call('POST', '/api/auth/signout')
    assert client.get(BASE + '/api/state', timeout=30).status_code == 401
    call('POST', '/api/auth/signin', json={'email': email, 'password': password})
    assert call('GET', '/api/state')['state']['assistants'][0]['draft']['name'] == 'Smoke Test Assistant'
    print('Real Qdrant indexing, cited retrieval, sign-out, sign-in and persisted workspace passed.', flush=True)


if __name__ == '__main__':
    main()
